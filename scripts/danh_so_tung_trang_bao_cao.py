"""Chân trang riêng với số thật 1..16, tránh trình xem lặp cached PAGE."""
from copy import deepcopy
from pathlib import Path
import shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'docs/bao-cao-ly-thuyet-phan-loai-van-ban.docx'
doc=Document(SOURCE)
before=[p.text for p in doc.paragraphs]
if len(doc.sections)==1:
    breaks=[p for p in doc.paragraphs if any(b.get(qn('w:type'))=='page' for b in p._p.iter(qn('w:br')))]
    assert len(breaks)==15, f'Expected 15 page breaks; found {len(breaks)}'
    prototype=deepcopy(doc.sections[-1]._sectPr)
    for p in breaks:
        for run in list(p._p.findall(qn('w:r'))): p._p.remove(run)
        sect=deepcopy(prototype)
        for tag in ['headerReference','footerReference','titlePg','pgNumType','type']:
            for element in list(sect.findall(qn('w:'+tag))): sect.remove(element)
        typ=OxmlElement('w:type');typ.set(qn('w:val'),'nextPage');sect.insert(0,typ)
        p._p.get_or_add_pPr().append(sect)
assert len(doc.sections)==16
for number,section in enumerate(doc.sections,1):
    section.different_first_page_header_footer=False
    section.footer_distance=Cm(1)
    # Mỗi section có footer khác, không dùng chung kết quả lưu tạm của PAGE.
    section.footer.is_linked_to_previous=False
    footer=section.footer
    for element in list(footer._element): footer._element.remove(element)
    p=footer.add_paragraph(str(number))
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(0)
    p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.line_spacing=1
    for run in p.runs:
        run.font.name='Times New Roman';run.font.size=Pt(13)
        run.font.color.rgb=RGBColor(0,0,0)
    numtype=section._sectPr.find(qn('w:pgNumType'))
    if numtype is None:
        numtype=OxmlElement('w:pgNumType');section._sectPr.append(numtype)
    numtype.set(qn('w:start'),str(number))
assert before==[p.text for p in doc.paragraphs]
assert [s.footer.paragraphs[0].text for s in doc.sections]==[str(i) for i in range(1,17)]
doc.save(SOURCE)
for name in ['bao-cao-tri-tue-nhan-tao-nhom-6.docx',
             'bao-cao-tri-tue-nhan-tao-nhom-6-co-so-trang.docx',
             'bao-cao-nhom-6-danh-so-1-den-16.docx']:
    shutil.copy2(SOURCE,ROOT/'docs'/name)
print('16 separate footers contain actual text:', ', '.join(str(i) for i in range(1,17)))

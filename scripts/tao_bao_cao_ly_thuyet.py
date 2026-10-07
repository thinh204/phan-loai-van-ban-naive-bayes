"""Tạo báo cáo lý thuyết bằng python-docx, không phụ thuộc ứng dụng demo."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/bao-cao-ly-thuyet-phan-loai-van-ban.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
sec.left_margin, sec.right_margin = Cm(3), Cm(2)
sec.different_first_page_header_footer = True
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.size = Pt(13)
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Times New Roman')
    fonts = st.element.get_or_add_rPr().get_or_add_rFonts()
    for attr in list(fonts.attrib):
        if 'Theme' in attr: del fonts.attrib[attr]
    for border in list(st.element.findall('.//' + qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(6)
normal.widow_control = True
for name, size in [('Title', 21), ('Heading 1', 17), ('Heading 2', 14)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = True
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(8)
    st.paragraph_format.keep_with_next = True

def p(text, bold=False, center=False):
    for old, new in {'N꜀ᵢ':'N(c,i)', 'N꜀ⱼ':'N(c,j)', 'n꜀':'n_c', 's꜀':'s_c', 'p꜀':'p_c', 'max꜀':'max_c', 'θcᵢ':'θ(c,i)'}.items():
        text = text.replace(old, new)
    x = doc.add_paragraph()
    x.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    x.add_run(text).bold = bold
    return x

def h(text, level=2):
    doc.add_paragraph(text, f'Heading {level}')

def eq(text):
    x = p(text, center=True)
    x.paragraph_format.space_before = Pt(3)
    x.paragraph_format.space_after = Pt(8)
    return x

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for c, width in zip(t.columns, widths): c.width = Cm(width)
    for i, val in enumerate(headers): t.rows[0].cells[i].text = val
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row): cells[i].text = str(val)
    borders = OxmlElement('w:tblBorders')
    for side in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement(f'w:{side}'); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
    t._tbl.tblPr.append(borders)
    for ri, row in enumerate(t.rows):
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci, cell in enumerate(row.cells):
            cell.width = Cm(widths[ci]); cell.vertical_alignment = 1
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e=OxmlElement(f'w:{side}');e.set(qn('w:w'),'65');e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            if ri == 0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7E6E6');pr.append(shade)
            for par in cell.paragraphs:
                par.paragraph_format.line_spacing = 1.05
                par.paragraph_format.space_after = Pt(2)
                par.paragraph_format.space_before = Pt(2)
                par.alignment = WD_ALIGN_PARAGRAPH.LEFT if ci == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in par.runs: run.font.size = Pt(12); run.bold = ri == 0
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def page(title):
    doc.add_page_break(); h(title, 1)

p('HỌC VIỆN CÔNG NGHỆ BƯU CHÍNH VIỄN THÔNG', True, True)
p('LỚP D23VHCN01-N', True, True)
doc.add_paragraph()
title = doc.add_paragraph('BÁO CÁO LÝ THUYẾT', 'Title'); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
p('MÔN HỌC  TRÍ TUỆ NHÂN TẠO', True, True)
doc.add_paragraph()
p('ĐỀ TÀI', True, True)
p('NGHIÊN CỨU VÀ TRÌNH BÀY PHƯƠNG PHÁP\nPHÂN LOẠI VĂN BẢN BẰNG\nMULTINOMIAL NAIVE BAYES', True, True)
doc.add_paragraph()
p('Nhóm sinh viên thực hiện', True, True)
table(['Họ và tên sinh viên','Mã số sinh viên'],[
    ['Lại Huy Thịnh','N23DVCN057'],['Nguyễn Trần Mạnh Dũng','N23DVCN01'],['Nguyễn Hữu Đức','N23DVCN012']], [10,6])
p('Giảng viên hướng dẫn: ...................................................', center=True)
doc.add_paragraph()
p('Tháng 10 năm 2026', center=True)

page('MỤC LỤC')
toc = [
('Lời mở đầu',3),('Chương 1 Tổng quan về phân loại văn bản',4),
('Chương 2 Tiền xử lý và biểu diễn văn bản',5),('2.1 Tiền xử lý văn bản',5),('2.2 Bag of Words và TF IDF',6),
('Chương 3 Cơ sở lý thuyết của Naive Bayes',7),('3.1 Định lý Bayes và giả định độc lập',7),('3.2 Mô hình Multinomial Naive Bayes',8),
('Chương 4 Quy trình học và dự đoán',9),('Chương 5 Ví dụ minh họa tính tay',10),('5.1 Dữ liệu và ước lượng tham số',10),('5.2 Tính điểm và giải thích dự đoán',11),
('Chương 6 Đánh giá phương pháp phân loại',12),('Chương 7 Nhận xét và hướng mở rộng',13),('7.1 Ưu điểm và hạn chế',13),('7.2 Ứng dụng và hướng phát triển',14),
('Kết luận',15),('Tài liệu tham khảo',16)]
for text, num in toc:
    x = doc.add_paragraph()
    x.paragraph_format.space_after = Pt(9)
    x.paragraph_format.tab_stops.add_tab_stop(Cm(15.5), WD_ALIGN_PARAGRAPH.RIGHT, 1)
    x.add_run(f'{text}\t{num}')

page('LỜI MỞ ĐẦU')
p('Văn bản xuất hiện trong thư điện tử, bài báo, thông báo, bình luận và nhiều nguồn thông tin khác. Khi số lượng tài liệu tăng, việc đọc và gán chủ đề cho từng tài liệu đòi hỏi nhiều thời gian. Phân loại văn bản tự động hỗ trợ tổ chức thông tin bằng cách gán văn bản vào các nhóm đã xác định trước. Đây là một bài toán phù hợp để tìm hiểu mối liên hệ giữa dữ liệu, biểu diễn đặc trưng và suy luận trong môn Trí tuệ nhân tạo.')
p('Báo cáo nghiên cứu phương pháp Multinomial Naive Bayes, một phương pháp phân loại xác suất thuộc nhóm học có giám sát. Nội dung đi từ khái niệm bài toán, cách chuyển văn bản thành véc-tơ, định lý Bayes và giả định độc lập có điều kiện đến quy tắc huấn luyện, dự đoán và đánh giá. Mục tiêu là làm rõ vì sao các tần suất từ có thể trở thành bằng chứng để lựa chọn nhãn cho một văn bản.')
h('Mục tiêu nghiên cứu')
p('Nhóm tập trung giải thích ý nghĩa các tham số của mô hình; trình bày cách ước lượng xác suất tiên nghiệm và xác suất từ theo lớp; phân tích vai trò của làm trơn và miền log; đồng thời minh họa đầy đủ phép tính bằng một tập dữ liệu nhỏ. Qua đó, người đọc có thể theo dõi một quyết định phân loại từ dữ liệu ban đầu đến kết quả cuối cùng.')
h('Phạm vi và phương pháp trình bày')
p('Phạm vi chính là phân loại đơn nhãn theo chủ đề: mỗi văn bản nhận một nhãn trong tập lớp cố định. Báo cáo sử dụng biểu diễn Bag of Words để giải thích chính xác mô hình đa thức và giới thiệu TF-IDF như một cách gán trọng số đặc trưng. Nội dung được xây dựng bằng tổng hợp tài liệu, diễn giải toán học, ví dụ tự xây dựng và phân tích các tình huống có thể gây sai lệch.')
p('Báo cáo gồm bảy chương, phần kết luận và tài liệu tham khảo. Ví dụ số trong chương 5 và chương 6 dùng để giải thích lý thuyết, không phải kết quả đo từ một hệ thống triển khai. Trọng tâm của đề tài là nguyên lý và giới hạn của phương pháp, để người học có cơ sở lựa chọn và sử dụng phương pháp phù hợp.')

page('CHƯƠNG 1 TỔNG QUAN VỀ PHÂN LOẠI VĂN BẢN')
h('1.1 Khái niệm và phát biểu bài toán')
p('Phân loại văn bản là xác định lớp của một tài liệu dựa trên nội dung của nó. Với tập lớp C = {c₁, c₂, …, cₖ}, hệ thống học một ánh xạ từ không gian biểu diễn văn bản vào C. Trong học có giám sát, dữ liệu học gồm các cặp (d, y), trong đó d là tài liệu và y là nhãn đã được cung cấp. Một tài liệu mới chưa biết nhãn được xử lý bằng cùng cách biểu diễn rồi đưa vào bộ phân loại [1].')
eq('D = {(d₁, y₁), …, (dₙ, yₙ)};    f(d) ∈ C')
p('Đầu vào có thể là nội dung thư, tiêu đề tin hoặc đoạn nhận xét. Đầu ra là nhãn như Thể thao, Công nghệ hay Thư rác. Đơn vị phân loại và định nghĩa từng nhãn cần được thống nhất trước khi gán nhãn dữ liệu. Chẳng hạn, một bài viết về ứng dụng công nghệ trong thi đấu có thể thuộc hai chủ đề; nếu bài toán chỉ cho phép một nhãn, nhóm cần quy tắc chọn chủ đề chính.')
h('1.2 Các dạng bài toán liên quan')
table(['Dạng bài toán','Đặc điểm đầu ra','Ví dụ'],[
['Nhị phân','Một trong hai lớp','Thư rác hoặc thư thường'],['Đa lớp đơn nhãn','Một lớp trong nhiều lớp','Thể thao hoặc Công nghệ hoặc Kinh tế'],['Đa nhãn','Có thể có nhiều nhãn','Một bài có cả nhãn Công nghệ và Giáo dục']], [4.2,5.6,6.2])
p('Phân loại khác với phân cụm: phân loại sử dụng tập lớp và nhãn học đã xác định; phân cụm tìm các nhóm theo mức tương đồng khi chưa có nhãn mục tiêu. Vì vậy, chất lượng nhãn học là một thành phần trực tiếp của bài toán phân loại.')
h('1.3 Khó khăn của dữ liệu văn bản')
p('Văn bản có độ dài khác nhau, từ đồng nghĩa, từ đa nghĩa, lỗi chính tả và những thông tin ngầm phụ thuộc ngữ cảnh. Số từ khác nhau thường lớn nhưng mỗi tài liệu chỉ chứa một phần nhỏ, tạo ra véc-tơ nhiều chiều và thưa. Với tiếng Việt, dấu cách thường phân tách âm tiết nên cần cân nhắc đơn vị từ; “máy tính” có thể cần được giữ như một đặc trưng thống nhất. Những đặc điểm này ảnh hưởng trực tiếp đến bước biểu diễn và các giả định của mô hình.')

page('CHƯƠNG 2 TIỀN XỬ LÝ VÀ BIỂU DIỄN VĂN BẢN')
h('2.1 Tiền xử lý văn bản')
p('Tiền xử lý nhằm đưa các văn bản về dạng nhất quán trước khi trích xuất đặc trưng. Các quyết định cần phù hợp với mục tiêu phân loại: thông tin có thể là nhiễu trong một bài toán nhưng lại hữu ích trong bài toán khác. Ví dụ, địa chỉ đường dẫn có thể được thay bằng ký hiệu chung khi phân loại chủ đề; trong nhận diện thư lừa đảo, đặc điểm của đường dẫn có thể là bằng chứng cần giữ.')
h('2.1.1 Chuẩn hóa và làm sạch')
p('Các bước thường được cân nhắc gồm chuẩn hóa Unicode, thống nhất khoảng trắng, xử lý thẻ HTML và chuyển chữ hoa thành chữ thường. Không nên xóa toàn bộ dấu câu, chữ số hoặc dấu tiếng Việt theo một quy tắc cố định. Nếu phân loại tài liệu về sản phẩm, “iPhone 15” và “iPhone 16” có thể cần được phân biệt. Chuẩn hóa phải bảo toàn thông tin liên quan đến nhãn.')
h('2.1.2 Tách đơn vị từ và xử lý từ phổ biến')
p('Token là đơn vị dùng để đếm hoặc gán trọng số. Token có thể là từ, cụm từ hoặc chuỗi ký tự. Với tiếng Việt, cách tách theo dấu cách tạo ra các âm tiết như “trí”, “tuệ”, “nhân”, “tạo”; một bộ tách từ có thể tạo các đơn vị “trí_tuệ” và “nhân_tạo”. Hai cách dẫn đến bộ từ vựng khác nhau và cần được dùng nhất quán ở cả quá trình học lẫn dự đoán.')
p('Từ dừng là những từ thường gặp mà người thiết kế cân nhắc loại bỏ. Tuy nhiên, từ phổ biến không đồng nghĩa với vô ích. Trong phân loại cảm xúc, bỏ từ “không” có thể làm “không tốt” trở nên giống “tốt”. Danh sách từ dừng cần được đánh giá theo bài toán thay vì áp dụng máy móc.')
h('2.1.3 Tính nhất quán của quy trình')
p('Giả sử dữ liệu học giữ cụm “máy_tính”, nhưng văn bản mới lại tách thành “máy” và “tính”. Khi đó các chiều đặc trưng không còn tương ứng, làm mất bằng chứng mà mô hình đã học. Vì vậy, bước chuẩn hóa, tách từ và ánh xạ từ vựng phải là cùng một quy trình. Trường hợp văn bản rỗng hoặc không có từ nào thuộc bộ từ vựng cũng cần được nhận biết, vì khi đó mô hình gần như chỉ dựa vào xác suất tiên nghiệm.')

page('2.2 BAG OF WORDS VÀ TF IDF')
h('2.2.1 Biểu diễn Bag of Words')
p('Bag of Words biểu diễn tài liệu bằng số lần xuất hiện của từng từ, không giữ thứ tự. Với bộ từ vựng V = [bóng, đội, máy, tính], câu “đội bóng bóng” có véc-tơ [2, 1, 0, 0]. Mỗi chiều có ý nghĩa cố định; giá trị bằng 0 biểu thị từ không xuất hiện. Véc-tơ được lưu ở dạng thưa để tránh lưu nhiều số 0 [3].')
eq('xᵢ = số lần token thứ i xuất hiện trong tài liệu d')
p('Hạn chế có thể thấy ngay từ biểu diễn: “đội thắng máy” và “máy thắng đội” có cùng số lần xuất hiện từng từ. Nếu nhãn phụ thuộc vai trò chủ thể và đối tượng, hai câu có ý nghĩa khác nhau nhưng mô hình không thể phân biệt bằng unigram. Có thể bổ sung bigram để giữ một phần thông tin cục bộ, với chi phí tăng kích thước từ vựng.')
h('2.2.2 Trọng số TF IDF')
p('TF-IDF kết hợp tần suất từ trong tài liệu và mức độ hiếm của từ trong tập tài liệu. Một lựa chọn công thức có làm trơn là [3]:')
eq('TF(t,d) = count(t,d)')
eq('IDF(t) = ln[(1 + N)/(1 + df(t))] + 1')
eq('w(t,d) = TF(t,d) × IDF(t)')
p('N là số tài liệu học; df(t) là số tài liệu học có chứa t, không phải tổng số lần t xuất hiện. Sau khi tính trọng số, có thể chuẩn hóa L2 bằng cách chia mỗi trọng số cho căn bậc hai của tổng bình phương các trọng số trong cùng tài liệu. Công thức TF và IDF có nhiều biến thể, vì vậy cần nêu rõ lựa chọn khi so sánh kết quả.')
p('Ví dụ tự xây dựng: có bốn tài liệu, từ “máy” xuất hiện trong ba tài liệu và từ “vi_mạch” xuất hiện trong một tài liệu. Khi đó IDF(máy) = ln(5/4) + 1 ≈ 1,2231; IDF(vi_mạch) = ln(5/2) + 1 ≈ 1,9163. Nếu mỗi từ xuất hiện một lần trong cùng tài liệu, “vi_mạch” có trọng số trước chuẩn hóa lớn hơn. Điều này thể hiện độ hiếm trong tập dữ liệu, không tự bảo đảm từ đó luôn hữu ích cho nhãn.')
h('2.2.3 Quan hệ với Multinomial Naive Bayes')
p('Mô hình đa thức được xây dựng cho số đếm không âm. TF-IDF tạo trọng số thực không âm và có thể dùng trong thực hành, nhưng khi đó cần hiểu đây là cách áp dụng quy tắc tính điểm có trọng số, thay vì coi mỗi trọng số là số lần rút token nguyên [4]. Bộ từ vựng và IDF chỉ được học từ dữ liệu huấn luyện để tránh đưa thông tin kiểm tra vào mô hình.')

page('CHƯƠNG 3 CƠ SỞ LÝ THUYẾT CỦA NAIVE BAYES')
h('3.1 Định lý Bayes và giả định độc lập')
p('Định lý Bayes liên hệ xác suất của một giả thuyết sau khi quan sát dữ liệu với xác suất của dữ liệu khi giả thuyết đúng. Trong phân loại, giả thuyết là “tài liệu thuộc lớp c”, còn bằng chứng là véc-tơ đặc trưng x. Với P(x) > 0, ta có:')
eq('P(c | x) = P(x | c) × P(c) / P(x)')
table(['Ký hiệu','Ý nghĩa'],[
['P(c)','Xác suất tiên nghiệm của lớp c'],['P(x | c)','Khả năng quan sát véc-tơ x khi biết lớp c'],['P(x)','Xác suất bằng chứng x'],['P(c | x)','Xác suất hậu nghiệm của lớp c khi quan sát x']], [4,12])
p('Bộ phân loại chọn lớp có hậu nghiệm lớn nhất, gọi là quyết định MAP. Vì cùng một x được xét cho mọi lớp, P(x) là hằng số trong phép so sánh và có thể bỏ khỏi quy tắc chọn lớp [2].')
eq('ĉ = argmax theo c thuộc C của P(x | c) × P(c)')
p('Khó khăn là việc mô tả đầy đủ phân phối của mọi tổ hợp đặc trưng. Naive Bayes dùng giả định độc lập có điều kiện: khi đã biết lớp, bằng chứng từ các đặc trưng được kết hợp theo dạng tích. Với mô hình token, ta giả sử mỗi vị trí token được sinh độc lập khi biết lớp và dùng cùng phân phối từ cho mọi vị trí [5].')
eq('P(t₁, …, tₘ | c) = ∏ⱼ P(tⱼ | c)')
p('“Độc lập có điều kiện” không có nghĩa là các từ độc lập trong mọi hoàn cảnh. Trong ngôn ngữ, từ “máy” thường đi cùng “tính”; mô hình vẫn đơn giản hóa quan hệ này để giảm số tham số. Cũng cần phân biệt các token tại từng vị trí với véc-tơ số đếm: khi tổng số token cố định, các số đếm của phân phối đa thức không độc lập với nhau. Giả định sinh token dẫn đến công thức đa thức ở phần tiếp theo.')

page('3.2 MÔ HÌNH MULTINOMIAL NAIVE BAYES')
h('3.2.1 Mô hình xác suất và quy tắc quyết định')
p('Giả sử từ vựng có V từ, tài liệu có véc-tơ đếm x = (x₁, …, xᵥ) và tổng m = Σᵢxᵢ. Với θcᵢ là xác suất từ i trong lớp c, θcᵢ ≥ 0 và Σᵢθcᵢ = 1, phân phối đa thức có dạng:')
eq('P(x | c, m) = [m! / (x₁! … xᵥ!)] × ∏ᵢ θcᵢ ^ xᵢ')
p('Trong quy tắc phân loại chuẩn, độ dài m được xem là đã biết và không mô hình hóa riêng theo lớp. Hệ số chứa giai thừa chỉ phụ thuộc x, không phụ thuộc c, nên được bỏ khi so sánh các lớp. Lấy log tự nhiên của phần còn lại cho điểm:')
eq('score(c,x) = ln P(c) + Σᵢ xᵢ ln θcᵢ')
eq('ĉ = argmax theo c thuộc C của score(c,x)')
p('Log là hàm tăng nên không làm thay đổi thứ tự các điểm xác suất dương. Việc cộng log thay cho nhân nhiều số nhỏ cũng giảm nguy cơ tràn dưới số thực. Điểm log thường âm; điểm ít âm hơn là điểm lớn hơn. Cần phân biệt điểm so sánh này với xác suất đã chuẩn hóa.')
h('3.2.2 Ước lượng tham số và làm trơn')
p('Gọi n là tổng số tài liệu học, n꜀ là số tài liệu lớp c, N꜀ᵢ là tổng số lần từ i xuất hiện trong các tài liệu của lớp c. Khi α > 0, dùng ước lượng:')
eq('P(c) = n꜀ / n')
eq('θcᵢ = (N꜀ᵢ + α) / (Σⱼ N꜀ⱼ + αV)')
p('α = 1 tương ứng làm trơn Laplace. Nếu một từ trong từ vựng chưa xuất hiện ở lớp c, làm trơn giữ xác suất của từ đó khác 0. Nếu không làm trơn, một số hạng bằng 0 có thể khiến toàn bộ tích của lớp bằng 0. α quá lớn kéo các phân phối từ về gần đều; α nhỏ giữ rõ hơn khác biệt số đếm. Không có một giá trị α tốt nhất cho mọi bài toán [2], [4].')
p('Làm trơn không tự thêm từ mới vào từ vựng. Từ chưa có trong từ vựng học thường bị bỏ qua khi biến đổi văn bản mới. Vì vậy, thiếu vốn từ và thiếu số đếm trong một lớp là hai vấn đề khác nhau.')

page('CHƯƠNG 4 QUY TRÌNH HỌC VÀ DỰ ĐOÁN')
h('4.1 Giai đoạn học')
p('Bước 1 là xác định tập lớp, thu thập tài liệu có nhãn và thống nhất quy tắc gán nhãn. Bước 2 là chia dữ liệu phục vụ huấn luyện và đánh giá trước khi học các thống kê từ vựng. Bước 3 là tiền xử lý dữ liệu học, xây dựng từ vựng và chuyển tài liệu thành véc-tơ. Bước 4 là cộng số đếm theo lớp, ước lượng P(c), θcᵢ và lưu các giá trị log. Nếu dùng TF-IDF, tổng trọng số thay thế tổng số đếm trong cách áp dụng thực hành.')
p('Kết quả của giai đoạn học gồm quy tắc tiền xử lý, ánh xạ từ vựng, các thống kê biểu diễn và tham số của bộ phân loại. Các thành phần phải được giữ cùng nhau; chỉ lưu bảng xác suất mà bỏ ánh xạ từ vựng sẽ khiến hệ thống không biết mỗi chiều đặc trưng tương ứng với từ nào.')
h('4.2 Giai đoạn dự đoán')
p('Văn bản mới được xử lý bằng đúng quy tắc đã học. Với mỗi lớp, bắt đầu từ ln P(c), rồi cộng xᵢ ln θcᵢ cho các đặc trưng có giá trị khác 0. Lớp có tổng điểm cao nhất được chọn. Nếu cần biểu diễn xác suất hậu nghiệm của mô hình, có thể chuẩn hóa các điểm s꜀ bằng công thức ổn định:')
eq('a = max꜀ s꜀;    p꜀ = exp(s꜀ − a) / Σⱼ exp(sⱼ − a)')
p('Các giá trị p꜀ cộng bằng 1 theo mô hình, nhưng mức 0,95 không tự chứng minh rằng 95% trường hợp tương tự sẽ đúng. Giả định đơn giản hóa có thể làm mô hình quá tự tin. Độ tin cậy của xác suất cần được đánh giá riêng nếu được dùng cho quyết định quan trọng [4].')
h('4.3 Chi phí tính toán và điều kiện áp dụng')
p('Gọi T là tổng token trong dữ liệu học, K là số lớp, V là số từ vựng. Với biểu diễn đếm, bước tích lũy số đếm có chi phí xấp xỉ O(T), và chuẩn bị bảng tham số có chi phí O(KV). Với một tài liệu có s đặc trưng khác 0, bước tính điểm có chi phí O(Ks), chưa kể tiền xử lý. Bảng tham số thường cần O(KV) bộ nhớ; lưu văn bản theo dạng thưa giúp giảm chi phí của ma trận tài liệu [2].')
p('Quy trình phù hợp khi tập lớp rõ ràng, có dữ liệu gán nhãn và đặc trưng không âm. Nếu véc-tơ đầu vào đã được biến đổi thành giá trị âm, không thể đưa trực tiếp vào mô hình đa thức như với số đếm từ thông thường.')

page('CHƯƠNG 5 VÍ DỤ MINH HỌA TÍNH TAY')
h('5.1 Dữ liệu và ước lượng tham số')
p('Nhóm xây dựng ví dụ gồm hai lớp Thể thao và Công nghệ, mỗi lớp hai tài liệu. Để tập trung vào phép tính, quy ước mỗi đơn vị cách nhau bởi dấu cách là một token, không loại từ nào và chỉ dùng Bag of Words. Đây là dữ liệu nhỏ minh họa, không đại diện cho ngôn ngữ tự nhiên đầy đủ.')
table(['Mã','Văn bản học','Lớp'],[
['d₁','bóng đội thắng','Thể thao'],['d₂','bóng bóng đội','Thể thao'],['d₃','máy tính nhanh','Công nghệ'],['d₄','máy máy tính','Công nghệ']], [1.6,8.8,5.6])
p('Từ vựng theo thứ tự cố định là [bóng, đội, thắng, máy, tính, nhanh], nên V = 6. Mỗi lớp có hai tài liệu và sáu token. Vì vậy P(Thể thao) = P(Công nghệ) = 2/4 = 1/2. Với làm trơn α = 1, mẫu số của mỗi lớp là 6 + 1 × 6 = 12.')
table(['Token','Đếm Thể thao','Đếm Công nghệ','θ Thể thao','θ Công nghệ'],[
['bóng',3,0,'4/12','1/12'],['đội',2,0,'3/12','1/12'],['thắng',1,0,'2/12','1/12'],['máy',0,3,'1/12','4/12'],['tính',0,2,'1/12','3/12'],['nhanh',0,1,'1/12','2/12']], [2.6,3.3,3.3,3.4,3.4])
p('Kiểm tra tính hợp lệ: với lớp Thể thao, tổng xác suất là (4 + 3 + 2 + 1 + 1 + 1)/12 = 1. Lớp Công nghệ cũng có tổng bằng 1. Những từ có số đếm 0 vẫn nhận xác suất 1/12. Các số này thể hiện ảnh hưởng của làm trơn, không có nghĩa là nhóm đã quan sát các từ đó trong lớp tương ứng.')
p('Đặt văn bản cần phân loại là “bóng đội máy”. Theo thứ tự từ vựng trên, véc-tơ đầu vào là x = [1, 1, 0, 1, 0, 0]. Ba từ có mặt đóng góp vào điểm dự đoán; các từ có số đếm bằng 0 không tạo thêm số hạng trong tổng điểm của mô hình đa thức.')

page('5.2 TÍNH ĐIỂM VÀ GIẢI THÍCH DỰ ĐOÁN')
h('5.2.1 So sánh hai lớp')
p('Bỏ hệ số đa thức chung cho hai lớp, điểm theo dạng tích của văn bản “bóng đội máy” là:')
eq('q(Thể thao) = (1/2) × (4/12) × (3/12) × (1/12) = 1/288')
eq('q(Công nghệ) = (1/2) × (1/12) × (1/12) × (4/12) = 1/864')
p('q là điểm chưa chuẩn hóa, không gọi trực tiếp là hậu nghiệm. Vì 1/288 lớn gấp ba lần 1/864, nhãn dự đoán là Thể thao. Nếu tính đầy đủ P(x | c,m), hệ số đa thức là 3!/(1!1!1!) = 6; hệ số này nhân vào cả hai điểm và không thay đổi kết quả so sánh.')
p('Ở miền log, điểm của lớp Thể thao xấp xỉ ln(1/288) = −5,6630; điểm của lớp Công nghệ xấp xỉ ln(1/864) = −6,7616. Chênh lệch 1,0986 bằng ln(3), phù hợp với tỷ số của hai điểm dạng tích.')
eq('P(Thể thao | x) = (1/288) / [(1/288) + (1/864)] = 0,75')
eq('P(Công nghệ | x) = 0,25')
h('5.2.2 Giải thích sự đóng góp của từ')
p('Từ “bóng” thiên về Thể thao với tỷ số xác suất 4; từ “đội” thiên về Thể thao với tỷ số 3; từ “máy” thiên về Công nghệ nên tỷ số Thể thao/Công nghệ là 1/4. Hai tiên nghiệm bằng nhau. Tích tỷ số là 4 × 3 × (1/4) = 3, nên bằng chứng tổng hợp nghiêng về Thể thao dù văn bản có từ “máy”.')
h('5.2.3 Những tình huống cần chú ý')
p('Nếu xét “bóng máy”, tỷ số bằng 4 × (1/4) = 1, nên hai lớp có điểm bằng nhau. Hệ thống cần quy tắc xử lý hòa; không nên diễn giải trường hợp này như một quyết định chắc chắn. Nếu xét “bóng bóng máy”, từ “bóng” được đếm hai lần nên tỷ số là 4² × (1/4) = 4. Điều đó minh họa ảnh hưởng của tần suất trong Multinomial Naive Bayes.')
p('Nếu văn bản chỉ có từ “robot” chưa nằm trong từ vựng, phép biến đổi hiện tại cho véc-tơ toàn 0. Điểm chỉ còn tiên nghiệm và hai lớp lại hòa. Muốn xử lý trường hợp này cần chính sách từ chưa biết hoặc yêu cầu thêm nội dung; làm trơn trên sáu từ đã có không tự giải quyết từ thứ bảy chưa được đưa vào biểu diễn.')

page('CHƯƠNG 6 ĐÁNH GIÁ PHƯƠNG PHÁP PHÂN LOẠI')
h('6.1 Nguyên tắc đánh giá')
p('Dữ liệu huấn luyện dùng để học tham số; dữ liệu xác thực dùng để lựa chọn cách biểu diễn và siêu tham số; dữ liệu kiểm tra được giữ lại để đánh giá cuối cùng. Khi ít dữ liệu, có thể dùng xác thực chéo trên phần dữ liệu học. Mỗi lần chia phải học lại từ vựng và IDF bên trong tập huấn luyện tương ứng. Học IDF từ toàn bộ dữ liệu trước khi chia làm cho thông tin kiểm tra đi vào quy trình [6], [7].')
h('6.2 Ma trận nhầm lẫn và các thước đo')
p('Ví dụ tự xây dựng: kiểm tra 20 văn bản với Thể thao được chọn làm lớp dương. Giả sử có ma trận sau; hàng là nhãn thực, cột là nhãn dự đoán:')
table(['Nhãn thực','Dự đoán Thể thao','Dự đoán Công nghệ'],[
['Thể thao',8,2],['Công nghệ',1,9]], [5.2,5.4,5.4])
p('Khi đó TP = 8, FN = 2, FP = 1, TN = 9. TP là dự đoán dương đúng; FN là mẫu dương bị bỏ sót; FP là mẫu âm bị gán dương; TN là dự đoán âm đúng. Từ ma trận có thể suy ra [8]:')
eq('Accuracy = (TP + TN)/(TP + TN + FP + FN) = 17/20 = 85%')
eq('Precision = TP/(TP + FP) = 8/9 ≈ 88,89%')
eq('Recall = TP/(TP + FN) = 8/10 = 80%')
eq('F1 = 2TP/(2TP + FP + FN) = 16/19 ≈ 84,21%')
p('Precision trả lời có bao nhiêu dự đoán Thể thao là đúng; Recall trả lời tìm được bao nhiêu văn bản Thể thao thực. Với đa lớp, tính các thước đo theo từng lớp. Macro-F1 là trung bình F1 của các lớp, mỗi lớp có trọng số bằng nhau; Weighted-F1 lấy trọng số theo số mẫu của lớp. Nếu mẫu số bằng 0, phải nêu rõ quy ước xử lý.')
h('6.3 Diễn giải kết quả có trách nhiệm')
p('Accuracy cao chưa đủ khi số mẫu các lớp chênh lệch lớn. Nếu 95% thư là thư thường, bộ phân loại luôn đoán thư thường đạt 95% Accuracy nhưng bỏ sót toàn bộ thư rác. Cần xem từng lớp, ma trận nhầm lẫn và loại sai lầm quan trọng đối với ứng dụng. Những tỷ lệ trong ví dụ trên chỉ minh họa cách tính, không phải chất lượng thực nghiệm của đề tài.')

page('CHƯƠNG 7 NHẬN XÉT VÀ HƯỚNG MỞ RỘNG')
h('7.1 Ưu điểm và hạn chế')
h('7.1.1 Ưu điểm')
p('Cấu trúc mô hình đơn giản giúp người học liên hệ được bảng số đếm, xác suất và quyết định phân loại. Có thể giải thích một dự đoán bằng các số hạng đóng góp của từ như trong chương 5. Quá trình học chủ yếu tích lũy thống kê nên phù hợp làm phương pháp cơ sở khi nghiên cứu bài toán mới. Khi dữ liệu được cập nhật, các thống kê số đếm cũng có thể được cộng dồn nếu giữ nhất quán tập lớp và bộ từ vựng.')
p('Dữ liệu văn bản thường có nhiều chiều nhưng thưa, trong khi tính điểm chỉ cần duyệt các chiều khác 0. Điều này giúp mô hình có chi phí dự đoán dễ kiểm soát. Lợi ích thực tế còn phụ thuộc số lớp, kích thước từ vựng, bộ tách từ và cách lưu dữ liệu; không nên quy đổi các đặc điểm lý thuyết thành một tốc độ chạy cụ thể khi chưa đo trên môi trường xác định.')
h('7.1.2 Hạn chế')
p('Giả định độc lập và biểu diễn túi từ làm mất quan hệ về trật tự và ngữ cảnh. Hai văn bản có cùng số đếm từ sẽ nhận cùng điểm dù cấu trúc câu khác nhau. Các từ liên quan mạnh có thể cung cấp bằng chứng bị tính lặp; xác suất hậu nghiệm vì vậy có thể quá tự tin. Từ đa nghĩa, cách diễn đạt mỉa mai và những nội dung cần tri thức ngoài văn bản đều là tình huống cần thận trọng.')
p('Chất lượng cũng bị giới hạn bởi nhãn và phạm vi dữ liệu học. Nếu phần lớn bài Thể thao được lấy từ một nguồn có cụm chữ cố định, mô hình có thể học dấu vết nguồn thay cho chủ đề. Khi nguồn mới không chứa dấu vết ấy, độ chính xác có thể giảm. Một nhãn hiếm hoặc được định nghĩa thiếu nhất quán cũng khiến các thống kê kém đại diện.')
h('7.1.3 Những nhận định cần tránh')
p('Không thể khẳng định Naive Bayes luôn tốt nhất, luôn cần ít dữ liệu hoặc hiểu ngữ nghĩa như con người. Tương tự, TF-IDF không bảo đảm luôn tốt hơn số đếm, và làm trơn không khôi phục thông tin đã bị mất do tiền xử lý. Các kết luận về chất lượng phải gắn với một bài toán, tập dữ liệu và cách đánh giá rõ ràng. Phân tích lý thuyết giúp dự đoán điểm mạnh và điểm yếu, còn kiểm chứng thực nghiệm là bước cần thiết nếu muốn đưa ra kết luận định lượng.')

page('7.2 ỨNG DỤNG VÀ HƯỚNG PHÁT TRIỂN')
h('7.2.1 Ứng dụng phù hợp')
p('Một ứng dụng là phân luồng thư điện tử vào nhóm hỗ trợ kỹ thuật, thanh toán hoặc yêu cầu chung. Khi xây dựng nhãn, cần quy định cách xử lý thư có nhiều yêu cầu, tránh buộc mô hình học các nhãn mâu thuẫn. Một ứng dụng khác là gợi ý chủ đề tin bài để người biên tập kiểm tra. Với các văn bản ngắn hoặc bằng chứng yếu, có thể cho phép từ chối dự đoán thay vì buộc chọn một nhãn có độ tin cậy thấp.')
p('Trong nhận diện thư rác, bỏ sót thư rác và chặn nhầm thư hợp lệ có hậu quả khác nhau. Vì vậy, cùng một mô hình xác suất có thể cần ngưỡng quyết định khác nhau tùy yêu cầu sử dụng. Ngưỡng phải được lựa chọn từ dữ liệu xác thực và đánh giá với chỉ tiêu phù hợp; không tự lấy một ngưỡng tùy ý rồi coi đó là bảo đảm an toàn.')
h('7.2.2 So sánh với một số hướng tiếp cận')
table(['Phương pháp','Cách khai thác văn bản','Điểm cần cân nhắc'],[
['Multinomial NB','Thống kê từ theo từng lớp','Giả định độc lập và mất thứ tự'],
['Logistic Regression','Học trực tiếp trọng số phân biệt lớp','Cần tối ưu tham số và điều chuẩn'],
['SVM tuyến tính','Học ranh giới với biên phân tách','Cần chọn tham số và biểu diễn phù hợp'],
['Mô hình ngữ cảnh','Biểu diễn phụ thuộc ngữ cảnh','Chi phí tài nguyên và thiết kế đánh giá']], [4,6,6])
p('Bảng so sánh mô tả sự khác nhau về cách tiếp cận, không xếp hạng chất lượng. Khi so sánh, cần giữ thống nhất dữ liệu, tập kiểm tra và chỉ số để tránh quy sự khác biệt của dữ liệu thành ưu thế của thuật toán.')
h('7.2.3 Hướng mở rộng')
p('Có thể nghiên cứu tách từ tiếng Việt, đặc trưng bigram hoặc n-gram ký tự để bổ sung thông tin cục bộ và giảm ảnh hưởng lỗi chính tả. Một hướng khác là lựa chọn α, giới hạn từ quá hiếm và đối chiếu Bag of Words với TF-IDF. Ngoài ra, cần phân tích các nhóm sai theo độ dài, tỷ lệ từ ngoài từ vựng và mức mất cân bằng nhãn. Mỗi thay đổi nên giải quyết một vấn đề cụ thể và được đánh giá trên dữ liệu chưa dùng để lựa chọn cấu hình.')

page('KẾT LUẬN')
p('Báo cáo đã trình bày phương pháp phân loại văn bản bằng Multinomial Naive Bayes theo tiến trình từ phát biểu bài toán đến quyết định dự đoán. Văn bản được chuyển thành đặc trưng số; các thống kê theo lớp dùng để ước lượng tiên nghiệm và xác suất từ; định lý Bayes cùng giả định sinh token độc lập tạo ra quy tắc chọn lớp. Làm trơn xử lý số đếm bằng 0 trong từ vựng, còn miền log giúp tính điểm thuận tiện và ổn định hơn.')
p('Ví dụ tính tay cho thấy một dự đoán là kết quả tổng hợp nhiều bằng chứng. Văn bản “bóng đội máy” được xếp vào Thể thao vì tỷ số bằng chứng tổng hợp là 3, dù có một từ thiên về Công nghệ. Ví dụ cũng chỉ ra trường hợp hòa điểm, vai trò của từ lặp và giới hạn khi gặp từ ngoài từ vựng. Những tình huống này giúp giải thích mô hình cụ thể hơn việc chỉ nêu một công thức.')
p('Đối với đánh giá, cần phân biệt dữ liệu học, dữ liệu xác thực và dữ liệu kiểm tra, đồng thời tránh học thống kê biểu diễn từ tập kiểm tra. Accuracy, Precision, Recall và F1 phản ánh các khía cạnh khác nhau; lựa chọn chỉ tiêu cần gắn với hậu quả của từng loại sai lầm. Một tỷ lệ dự đoán cao không đủ để kết luận phương pháp phù hợp với mọi nguồn văn bản hoặc mọi ngôn ngữ.')
h('Kết quả đạt được về mặt lý thuyết')
p('Nhóm hệ thống hóa được quan hệ giữa tiền xử lý, biểu diễn véc-tơ và suy luận xác suất; trình bày đầy đủ các thành phần của Multinomial Naive Bayes; xây dựng ví dụ số có thể kiểm tra lại; và phân tích điều kiện áp dụng, ưu điểm, hạn chế. Nội dung đáp ứng mục tiêu nghiên cứu và trình bày một phương pháp phân loại văn bản trong môn Trí tuệ nhân tạo.')
h('Định hướng nghiên cứu tiếp theo')
p('Khi mở rộng đề tài, cần ưu tiên định nghĩa nhãn rõ ràng và lựa chọn biểu diễn phù hợp với tiếng Việt trước khi tăng độ phức tạp thuật toán. Multinomial Naive Bayes có thể dùng làm phương pháp cơ sở để đối chiếu các hướng tiếp cận khác trong cùng điều kiện đánh giá. Việc mở rộng chỉ có ý nghĩa khi giải thích được vấn đề cần khắc phục và đưa ra bằng chứng kiểm chứng phù hợp.')

page('TÀI LIỆU THAM KHẢO')
refs = [
('[1]','Manning, C. D., Raghavan, P. và Schütze, H. (2008). Introduction to Information Retrieval. Cambridge University Press. Chương 13, The text classification problem.','https://nlp.stanford.edu/IR-book/html/htmledition/the-text-classification-problem-1.html'),
('[2]','Manning, C. D., Raghavan, P. và Schütze, H. (2008). Naive Bayes text classification. Introduction to Information Retrieval, mục 13.2.','https://nlp.stanford.edu/IR-book/html/htmledition/naive-bayes-text-classification-1.html'),
('[3]','scikit-learn developers. Feature extraction. Tài liệu về Bag of Words, token hóa và TF-IDF.','https://scikit-learn.org/stable/modules/feature_extraction.html'),
('[4]','scikit-learn developers. Naive Bayes. Tài liệu về mô hình đa thức, làm trơn và lưu ý về ước lượng xác suất.','https://scikit-learn.org/stable/modules/naive_bayes.html'),
('[5]','Manning, C. D., Raghavan, P. và Schütze, H. (2008). Properties of Naive Bayes. Introduction to Information Retrieval, mục 13.4.','https://nlp.stanford.edu/IR-book/html/htmledition/properties-of-naive-bayes-1.html'),
('[6]','scikit-learn developers. Common pitfalls and recommended practices. Mục về rò rỉ dữ liệu và tính nhất quán của tiền xử lý.','https://scikit-learn.org/stable/common_pitfalls.html'),
('[7]','scikit-learn developers. Cross-validation evaluating estimator performance.','https://scikit-learn.org/stable/modules/cross_validation.html'),
('[8]','scikit-learn developers. Metrics and scoring quantifying the quality of predictions. Mục về đánh giá phân loại.','https://scikit-learn.org/stable/modules/model_evaluation.html')]
for num, title, url in refs:
    x = p(f'{num} {title}')
    x.paragraph_format.keep_with_next = True
    x = p(url)
    x.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in x.runs: r.font.size = Pt(10)
p('Ngày truy cập tài liệu trực tuyến: 07/10/2026.')

foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
doc.core_properties.title='Nghiên cứu và trình bày phương pháp phân loại văn bản bằng Multinomial Naive Bayes'
doc.core_properties.subject='Báo cáo lý thuyết môn Trí tuệ nhân tạo'
doc.core_properties.author='Lại Huy Thịnh; Nguyễn Trần Mạnh Dũng; Nguyễn Hữu Đức'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)

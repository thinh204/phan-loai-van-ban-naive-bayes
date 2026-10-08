"""Bổ sung mô tả nghiên cứu và thiết kế, không thay đổi ứng dụng."""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/bao-cao-tri-tue-nhan-tao-nhom-6-hoan-thien.docx'
doc = Document(ROOT / 'docs/bao-cao-nhom-6-danh-so-1-den-16.docx')
body = doc._element.body
prototype = deepcopy(doc.sections[-1]._sectPr)
pages, current = [], []
for element in list(body):
    if element.tag == qn('w:sectPr'):
        continue
    element = deepcopy(element)
    sects = list(element.iter(qn('w:sectPr')))
    if sects:
        for sect in sects:
            sect.getparent().remove(sect)
        if ''.join(element.itertext()).strip():
            current.append(element)
        pages.append(current)
        current = []
    else:
        current.append(element)
pages.append(current)
assert len(pages) == 16
for element in list(body):
    if element.tag != qn('w:sectPr'):
        body.remove(element)

def p(text, bold=False):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    par.paragraph_format.line_spacing = 1.16
    par.paragraph_format.space_after = Pt(7)
    run = par.add_run(text)
    run.bold = bold
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0,0,0)
    return par

def h(text, level=2):
    par = doc.add_paragraph(text, 'Heading '+str(level))
    par.paragraph_format.space_before = Pt(8)
    par.paragraph_format.space_after = Pt(8)
    return par

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.autofit=False
    for row in [headers] + rows:
        cells = t.rows[0].cells if row is headers else t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value); cells[i].width=Cm(widths[i])
            for par in cells[i].paragraphs:
                par.paragraph_format.space_after=Pt(5)
                par.paragraph_format.space_before=Pt(5)
                par.paragraph_format.line_spacing=1.05
                for r in par.runs:
                    r.font.size=Pt(12); r.bold=row is headers
                    r.font.color.rgb=RGBColor(0,0,0)
            if row is headers:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');cells[i]._tc.get_or_add_tcPr().append(sh)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    return t

def boundary():
    par=doc.add_paragraph()
    par.paragraph_format.space_after=Pt(0);par.paragraph_format.space_before=Pt(0)
    par.paragraph_format.line_spacing=Pt(1)
    sect=deepcopy(prototype)
    for tag in ['headerReference','footerReference','titlePg','pgNumType','type']:
        for e in list(sect.findall(qn('w:'+tag))):sect.remove(e)
    typ=OxmlElement('w:type');typ.set(qn('w:val'),'nextPage');sect.insert(0,typ)
    par._p.get_or_add_pPr().append(sect)

def old(index):
    for element in pages[index]:
        body.insert(len(body)-1, deepcopy(element))

old(0);boundary()
h('MỤC LỤC',1).alignment=WD_ALIGN_PARAGRAPH.CENTER
entries=[('LỜI MỞ ĐẦU',3),('CHƯƠNG 1 TỔNG QUAN VỀ PHÂN LOẠI VĂN BẢN',4),('CHƯƠNG 2 TIỀN XỬ LÝ VÀ BIỂU DIỄN VĂN BẢN',5),('2.2 Bag of Words và TF IDF',6),('CHƯƠNG 3 CƠ SỞ LÝ THUYẾT CỦA NAIVE BAYES',7),('3.2 Mô hình Multinomial Naive Bayes',8),('CHƯƠNG 4 QUY TRÌNH HỌC VÀ DỰ ĐOÁN',9),('CHƯƠNG 5 VÍ DỤ MINH HỌA TÍNH TAY',10),('5.2 Tính toán và suy luận',11),('CHƯƠNG 6 ĐÁNH GIÁ PHƯƠNG PHÁP',12),('CHƯƠNG 7 PHÂN TÍCH VÀ THẢO LUẬN',13),('7.3 Ứng dụng và hướng mở rộng',14),('CHƯƠNG 8 Ý TƯỞNG VÀ CHỨC NĂNG HỆ THỐNG',15),('8.3 Yêu cầu và các tình huống sử dụng',16),('CHƯƠNG 9 DỮ LIỆU VÀ LOGIC XỬ LÝ',17),('9.3 Huấn luyện và lựa chọn mô hình',18),('CHƯƠNG 10 MÔ TẢ GIAO DIỆN VÀ KẾT QUẢ',19),('10.3 Giải thích và lịch sử dự đoán',20),('CHƯƠNG 11 THƯ VIỆN VÀ ĐỊNH HƯỚNG ĐÁNH GIÁ',21),('11.3 Đánh giá và giới hạn triển khai',22),('KẾT LUẬN',23),('TÀI LIỆU THAM KHẢO',24)]
pdfmetrics.registerFont(TTFont('TNR','C:/Windows/Fonts/times.ttf'))
pdfmetrics.registerFont(TTFont('TNRB','C:/Windows/Fonts/timesbd.ttf'))
corrections={
    '5.2 Tính toán và suy luận':'5.2 Tính điểm và giải thích dự đoán',
    'CHƯƠNG 6 ĐÁNH GIÁ PHƯƠNG PHÁP':'CHƯƠNG 6 ĐÁNH GIÁ PHƯƠNG PHÁP PHÂN LOẠI',
    'CHƯƠNG 7 PHÂN TÍCH VÀ THẢO LUẬN':'CHƯƠNG 7 NHẬN XÉT VÀ HƯỚNG MỞ RỘNG',
    '7.3 Ứng dụng và hướng mở rộng':'7.2 Ứng dụng và hướng phát triển',
}
entries=[(corrections.get(label,label),num) for label,num in entries]
t=doc.add_table(rows=0,cols=2);t.autofit=False
t.columns[0].width=Cm(15.3);t.columns[1].width=Cm(.7)
for label,num in entries:
    cells=t.add_row().cells;cells[0].width=Cm(15.3);cells[1].width=Cm(.7)
    font=12
    available=15.1/2.54*72
    major=label.startswith(('CHƯƠNG','KẾT','LỜI','TÀI'))
    metric='TNRB' if major else 'TNR'
    dots=max(3,int((available-pdfmetrics.stringWidth(label+'  ',metric,font))/pdfmetrics.stringWidth('.',metric,font)))
    cells[0].text=label+' '+'.'*dots;cells[1].text=str(num)
    for i,c in enumerate(cells):
        margins=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');margins.append(e)
        c._tc.get_or_add_tcPr().append(margins)
        for par in c.paragraphs:
            par.paragraph_format.left_indent=Cm(0)
            par.paragraph_format.right_indent=Cm(0)
            par.paragraph_format.first_line_indent=Cm(0)
            par.paragraph_format.line_spacing=Pt(18)
            par.paragraph_format.space_after=Pt(9)
            par.alignment=WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.LEFT
            for r in par.runs:r.font.size=Pt(font);r.bold=major
boundary()
for index in range(2,14):
    old(index);boundary()

h('CHƯƠNG 8 Ý TƯỞNG VÀ CHỨC NĂNG HỆ THỐNG',1)
h('8.1 Ý tưởng và phạm vi giải quyết')
p('Nhóm đề xuất hệ thống hỗ trợ xác định chủ đề của một đoạn văn bản tiếng Anh. Người dùng nhập nội dung, hệ thống biểu diễn bằng TF-IDF và áp dụng Multinomial Naive Bayes để chọn một trong bốn chủ đề đã học. Kết quả kèm phân bố xác suất, từ đặc trưng và cảnh báo giúp người dùng xem xét quyết định. Đây là bài toán học có giám sát, phân loại đa lớp nhưng đơn nhãn; hệ thống không tự sinh thêm chủ đề.')
p('Bốn lớp tương ứng với đồ họa máy tính, bóng chày, khoa học vũ trụ và chính trị tổng hợp. Nhãn tiếng Việt chỉ phục vụ hiển thị; dữ liệu và nội dung đầu vào của mô hình hiện tại là tiếng Anh. Các ví dụ tiếng Việt ở chương 2, 5 và 6 dùng để diễn giải lý thuyết và phép tính, không mô tả dữ liệu đã huấn luyện hệ thống.')
h('8.2 Các chức năng cần thiết')
table(['Chức năng','Mục đích và kết quả'],[
['Nhập văn bản và chọn mẫu','Nhập nội dung tiếng Anh hoặc chọn ví dụ có sẵn để khảo sát.'],
['Phân loại chủ đề','Trả về một nhãn trong bốn lớp và xác suất theo mô hình.'],
['Giải thích dự đoán','Hiển thị trọng số TF-IDF và bằng chứng từ các đặc trưng.'],
['Cảnh báo dữ liệu yếu','Thông báo khi nội dung rỗng, quá ngắn, ít từ thuộc từ vựng hoặc xác suất thấp.'],
['Lịch sử trong phiên','Xem các lần dự đoán, tải CSV hoặc xóa lịch sử của phiên hiện tại.'],
['Thông tin nghiên cứu','Xem nguyên lý, chỉ số đánh giá và giới hạn sử dụng.']],[4.2,11.8])
p('Người học là tác nhân sử dụng giao diện; nhóm nghiên cứu chuẩn bị dữ liệu và mô hình ở giai đoạn riêng. Giao diện hiện tại không cung cấp đăng nhập, quản trị người dùng, tải tài liệu Word/PDF hay huấn luyện trực tiếp trên web.')
boundary()
h('8.3 Yêu cầu và các tình huống sử dụng')
p('Tình huống chính: người dùng mở trang, đọc phạm vi tiếng Anh và bốn chủ đề, nhập hoặc chọn một văn bản mẫu rồi nhấn Phân loại. Hệ thống kiểm tra dữ liệu, dùng bộ biến đổi và mô hình đã học, sau đó hiển thị nhãn, xác suất, thời gian xử lý và phần giải thích. Một lần dự đoán với nội dung không rỗng được bổ sung vào lịch sử phiên, kể cả khi có cảnh báo.')
p('Tình huống dữ liệu rỗng: nội dung chỉ gồm khoảng trắng được coi là rỗng sau khi loại khoảng trắng ở hai đầu. Giao diện thông báo cần nhập văn bản, không gọi chức năng phân loại và không ghi một kết quả mới vào lịch sử. Tình huống ngoài từ vựng: hệ thống vẫn có thể chọn một lớp từ xác suất tiên nghiệm, nhưng phải cảnh báo rằng không có bằng chứng từ nội dung để người dùng tránh hiểu nhầm.')
h('8.4 Yêu cầu chất lượng')
p('Tính nhất quán: dữ liệu học và văn bản mới phải dùng cùng bộ từ vựng, IDF, thứ tự lớp và cách trích xuất đặc trưng. Tính tái lập: lưu tham số, hạt ngẫu nhiên, cách chia dữ liệu và phiên bản thư viện. Khả năng giải thích: nêu rõ ý nghĩa của xác suất và từ hỗ trợ; không diễn đạt xác suất cao thành bảo đảm đúng.')
p('Khả năng sử dụng: tên lớp, nút bấm, cảnh báo và bảng kết quả phải dễ đọc; giữ nhãn kỹ thuật để đối chiếu dữ liệu gốc. Khả năng vận hành: nạp mô hình một lần và tái sử dụng cho các yêu cầu, xử lý thiếu tệp mô hình bằng thông báo rõ ràng. Lịch sử là trạng thái trong phiên Streamlit, không phải cơ sở dữ liệu lưu vĩnh viễn; kết nối hoặc tải lại phiên có thể làm mất trạng thái.')
h('8.5 Kiến trúc khái niệm')
p('Giai đoạn nghiên cứu: dữ liệu có nhãn → kiểm tra dữ liệu → học TF-IDF trong từng tập huấn luyện → chọn alpha bằng kiểm định chéo → học mô hình cuối → đánh giá tập kiểm tra → lưu hiện vật mô hình. Giai đoạn sử dụng: giao diện Streamlit → dịch vụ phân loại → TF-IDF đã học → Multinomial Naive Bayes đã học → kết quả và giải thích → giao diện.')
p('Các tệp mô hình, bộ TF-IDF và danh sách lớp phải thuộc cùng một lần huấn luyện. Tệp kết quả đánh giá là dữ liệu tham chiếu cho trang giới thiệu; thao tác nhập một văn bản không làm huấn luyện lại mô hình.')
boundary()
h('CHƯƠNG 9 DỮ LIỆU VÀ LOGIC XỬ LÝ',1)
h('9.1 Dữ liệu phù hợp với phương pháp')
p('Hệ thống sử dụng bốn nhóm của bộ 20 Newsgroups, gồm 3.729 văn bản: 2.239 mẫu huấn luyện và 1.490 mẫu kiểm tra. Đây là một tập con của bộ dữ liệu gốc, không phải toàn bộ 20 chủ đề. Mỗi văn bản có một nhãn chủ đề; cách biểu diễn thưa, không âm phù hợp với quy tắc tính điểm của Multinomial Naive Bayes dùng TF-IDF [4, 9].')
table(['Nhãn gốc','Ý nghĩa hiển thị','Mẫu kiểm tra'],[
['comp.graphics','Đồ họa máy tính',389],['rec.sport.baseball','Thể thao – Bóng chày',397],['sci.space','Khoa học vũ trụ',394],['talk.politics.misc','Chính trị tổng hợp',310]],[5,8,3])
p('Lược đồ dữ liệu gồm text là nội dung, target là mã lớp, target_name là tên nhóm và split là train hoặc test. Mã lớp phải được ánh xạ bằng danh sách tên lớp của bộ dữ liệu, không suy đoán từ thứ tự hiển thị. Chẳng hạn, một bài thảo luận về quỹ đạo vệ tinh có thể mang target_name là sci.space; nhãn thật do dữ liệu cung cấp, không do mô hình tự tạo khi học.')
h('9.2 Kiểm tra chất lượng và tiền xử lý')
p('Khi tải dữ liệu, chương trình bỏ headers, footers và quotes nhằm giảm dấu hiệu từ nguồn thư và nội dung trích dẫn. Chương trình thống kê phân bố lớp, giá trị thiếu, văn bản rỗng và nội dung trùng; việc thống kê không đồng nghĩa đã tự động xóa tất cả các trường hợp này. Khi nghiên cứu mở rộng, cần quyết định xử lý từng nhóm và ghi lại ảnh hưởng đến tập dữ liệu.')
p('Dữ liệu dùng cách chia train/test có sẵn của bộ nguồn. Không mô tả thí nghiệm hiện tại là chia ngẫu nhiên 75/25. Dữ liệu kiểm tra không được dùng để học bộ từ vựng, IDF hoặc lựa chọn alpha. Nguồn thảo luận trực tuyến có thể chứa văn phong cũ và lệch miền so với tin tức hoặc văn bản người dùng nhập hôm nay.')
boundary()
h('9.3 Huấn luyện và lựa chọn mô hình')
p('Bộ TF-IDF chuyển chữ về thường, dùng unigram, min_df = 2 và max_df = 0,95. Một từ phải xuất hiện trong ít nhất hai tài liệu học và không xuất hiện trong hơn 95% tài liệu học để được giữ. Các thiết lập còn lại theo TfidfVectorizer: số đếm TF thông thường, IDF có làm trơn và chuẩn hóa L2. Hệ thống hiện có 13.068 đặc trưng. Không áp dụng riêng bộ tách từ tiếng Việt, stemming, lemmatization hoặc danh sách từ dừng [3].')
p('Nhóm khảo sát alpha thuộc {0,01; 0,05; 0,1; 0,2; 0,5; 1; 1,5; 2} bằng StratifiedKFold 5 phần, xáo trộn với random_state = 42. Trong mỗi phần, Pipeline học lại TF-IDF chỉ trên phần huấn luyện rồi đánh giá phần xác thực. Chọn alpha có F1 macro trung bình cao nhất, sau đó học lại trên toàn bộ tập huấn luyện. Mô hình được lưu hiện tại dùng alpha = 0,1; alpha = 1 là làm trơn Laplace, còn 0,1 là làm trơn cộng với mức nhỏ hơn.')
h('9.4 Logic phân loại một văn bản mới')
p('Bước 1: loại khoảng trắng ở hai đầu và chặn nội dung rỗng ở giao diện. Bước 2: dùng transform của bộ TF-IDF đã học, không gọi fit trên đầu vào mới. Bước 3: tính điểm log của từng lớp theo chương 3, chọn lớp lớn nhất và chuẩn hóa thành phân bố xác suất. Bước 4: ánh xạ nhãn, trích đặc trưng khác 0, tạo cảnh báo và trả kết quả cho giao diện.')
p('Nội dung ngắn hơn 10 ký tự hoặc dưới 3 từ được cảnh báo ngắn. Nếu số đặc trưng khác 0 bằng 0 thì cảnh báo ngoài từ vựng; từ 1 đến 2 thì cảnh báo ít bằng chứng. Số này đếm các đặc trưng phân biệt có trọng số khác 0, không phải tổng số từ xuất hiện. Xác suất lớn nhất dưới 0,60 tạo cảnh báo độ tin cậy thấp. Cờ chưa chắc chắn được bật khi rỗng, ngoài từ vựng, ít đặc trưng hoặc xác suất thấp; riêng cảnh báo ngắn không tự bật cờ này.')
p('Các ngưỡng trên là quy tắc hỗ trợ giao diện hiện tại, không phải định lý đảm bảo nhận biết mọi văn bản sai chủ đề. Văn bản ngoài bốn nhóm vẫn có thể nhận xác suất cao; cần diễn giải trong phạm vi tập lớp đã học.')
boundary()
h('CHƯƠNG 10 MÔ TẢ GIAO DIỆN VÀ KẾT QUẢ',1)
h('10.1 Bố cục trang chính')
p('Giao diện web dùng Streamlit, gồm thanh bên và vùng nội dung chính. Thanh bên trình bày phiên bản, thuật toán, alpha, kích thước từ vựng, bốn nhãn, thông tin đánh giá và danh sách văn bản mẫu. Vùng chính đặt tên hệ thống phía trên, tiếp theo là bốn thẻ nội dung. Mô tả này tương ứng với app.py của dự án; đây là thiết kế giao diện và luồng tương tác, không phải hướng dẫn viết chương trình.')
table(['Vùng giao diện','Nội dung và thao tác'],[
['Thanh bên','Thông tin mô hình và lựa chọn văn bản mẫu tiếng Anh.'],
['Dự Đoán & Giải Thích','Ô nhập nội dung, nút Phân loại, cảnh báo, nhãn và phần giải thích.'],
['Giới Thiệu Mô Hình','Nguyên lý TF-IDF, Naive Bayes và kiến trúc xử lý.'],
['Đánh Giá Thực Nghiệm','Các chỉ số, kết quả theo lớp, ma trận nhầm lẫn và bảng khảo sát alpha.'],
['Giới Hạn & Lưu Ý','Phạm vi ngôn ngữ, dữ liệu yếu và giới hạn của mô hình.']],[5.7,10.3])
h('10.2 Luồng tương tác trên thẻ dự đoán')
p('Ô Nội dung văn bản (tiếng Anh) cho phép nhập đoạn nhiều dòng. Người dùng có thể chọn ví dụ ở thanh bên trước khi nhấn Phân loại. Khi có kết quả, phía trên hiển thị nhãn dự đoán bằng tên dễ đọc, xác suất lớn nhất và độ trễ tính bằng mili giây. Cảnh báo được đặt cùng khu vực kết quả để người dùng đọc trước khi sử dụng nhãn.')
p('Phía dưới là bảng xác suất của đủ bốn lớp và biểu đồ so sánh, tiếp theo là các từ có trọng số và phần giải thích. Cuối khu vực dự đoán là lịch sử trong phiên, nút tải CSV và thao tác xóa lịch sử. Không có bước chọn nhãn đúng bắt buộc hoặc chức năng sửa nhãn để mô hình học trực tiếp từ người dùng.')
p('Khi trình bày trên lớp, có thể minh họa lần lượt một đoạn đúng chủ đề, một đoạn quá ngắn và một đoạn không có từ thuộc từ vựng. Mục đích là giải thích cách giao diện phản ánh bằng chứng của mô hình; không gán trước kết quả hay xác suất nếu chưa có đầu ra tương ứng.')
boundary()
h('10.3 Giải thích và lịch sử dự đoán')
p('Bảng đặc trưng cho biết token và trọng số TF-IDF trong văn bản hiện tại. Trọng số cao phản ánh mức nổi bật theo cách biểu diễn, chưa đủ để kết luận token đó ủng hộ lớp dự đoán. Vì vậy, hệ thống bổ sung điểm hỗ trợ dựa trên xác suất từ theo lớp đã học.')
p('Với lớp dự đoán ĉ, điểm hỗ trợ của đặc trưng i là:')
p('gᵢ = xᵢ × [ln P(wᵢ | ĉ) − max theo k ≠ ĉ của ln P(wᵢ | k)]',True)
p('Điểm dương nghĩa là đặc trưng nghiêng về lớp dự đoán so với đối thủ có xác suất từ lớn nhất tại đặc trưng đó; điểm âm nghiêng về một lớp khác. Đối thủ có thể khác nhau giữa các từ. Vì công thức không cộng xác suất tiên nghiệm và không dùng một đối thủ cố định cho mọi từ, không được xem tổng các điểm là toàn bộ chênh lệch điểm phân loại hay bằng chứng nhân quả.')
p('Giao diện phân mức hỗ trợ: trên 0,25 là rất mạnh; trên 0,05 đến 0,25 là ủng hộ tích cực; từ −0,05 đến 0,05 là trung tính; nhỏ hơn −0,05 là hơi thiên lớp khác. Các mức này giúp đọc kết quả, không thay thế đánh giá độ chính xác.')
h('10.4 Thông tin lưu trong phiên và diễn giải xác suất')
p('Mỗi bản ghi lịch sử gồm thời điểm, nội dung xem trước tối đa 80 ký tự, nhãn dự đoán, xác suất phần trăm, số đặc trưng thuộc từ vựng, cờ chưa chắc chắn và độ trễ. CSV dùng UTF-8 có dấu BOM để thuận tiện đọc tiếng Việt. Xóa lịch sử chỉ tác động trạng thái phiên hiện tại, không xóa dữ liệu huấn luyện hoặc tệp mô hình.')
p('Xác suất dự đoán là hậu nghiệm theo giả định Naive Bayes và các tham số đã học. Giá trị 90% không có nghĩa đã kiểm chứng chắc chắn rằng 90 trong mọi 100 văn bản tương tự sẽ đúng; muốn diễn giải như vậy phải khảo sát hiệu chuẩn xác suất. Khi dữ liệu ngắn, đa chủ đề, ngoài miền hoặc ngoài từ vựng, cần đọc cảnh báo và xem lại nội dung trước khi chấp nhận nhãn.')
boundary()
h('CHƯƠNG 11 THƯ VIỆN VÀ ĐỊNH HƯỚNG ĐÁNH GIÁ',1)
h('11.1 Thư viện có thể triển khai phương pháp')
table(['Thành phần','Vai trò trong thiết kế'],[
['Python','Ngôn ngữ tổ chức dữ liệu, quy trình học và dịch vụ phân loại.'],
['Pandas','DataFrame cho text, target, target_name, split; thống kê dữ liệu và xuất lịch sử CSV [10].'],
['scikit-learn','Tải 20 Newsgroups; TF-IDF; MultinomialNB; Pipeline; StratifiedKFold; các chỉ số và ma trận nhầm lẫn [3, 4, 9].'],
['NumPy và SciPy','Mảng số, phép toán và biểu diễn ma trận đặc trưng thưa; giảm lưu trữ các phần tử bằng 0.'],
['joblib','Lưu và nạp lại mô hình, bộ TF-IDF và danh sách lớp từ các tệp tin cậy [12].'],
['Streamlit','Ô nhập, nút bấm, thẻ nội dung, bảng, biểu đồ, bộ nhớ đệm mô hình và lịch sử phiên [11].'],
['pytest','Hỗ trợ kiểm tra tính đúng của các chức năng; không tham gia công thức phân loại.']],[4,12])
p('requirements.txt của dự án cố định scikit-learn ở phiên bản 1.7.2; các thư viện còn lại được khai báo theo mức phiên bản tối thiểu. Việc nghiên cứu tài liệu mới cần phân biệt với phiên bản thực tế của môi trường và tệp mô hình đã lưu. Không thể giả định mọi phiên bản thư viện đều nạp mô hình cũ giống nhau.')
h('11.2 Các vấn đề liên quan cần nghiên cứu')
p('Bộ dữ liệu và cách chia ảnh hưởng khả năng khái quát; từ vựng và IDF ảnh hưởng biểu diễn; alpha điều chỉnh mức làm trơn; sự mất cân bằng lớp ảnh hưởng cách đọc chỉ số. Vì thế cần khảo sát từng yếu tố trên phần huấn luyện/xác thực, giữ quy trình cố định khi so sánh và ghi rõ các tham số. Các hướng bigram, Complement Naive Bayes hoặc mô hình tuyến tính chỉ là phương án mở rộng; hệ thống hiện tại vẫn dùng unigram và Multinomial Naive Bayes.')
boundary()
h('11.3 Đánh giá và giới hạn triển khai')
p('Kế hoạch đánh giá sử dụng Accuracy, Precision, Recall và F1 theo chương 6. F1 macro lấy trung bình các lớp nên giúp quan sát đồng đều bốn chủ đề; F1 weighted có trọng số theo số mẫu. Ma trận nhầm lẫn chỉ ra cặp lớp dễ nhầm để chọn ví dụ phân tích. Không chỉ chọn Accuracy cao rồi bỏ qua lớp có Recall thấp.')
p('Để đối chiếu thiết kế với hệ thống đã có, bảng sau trích số liệu lưu trong results/evaluation_summary.json. Đây là số liệu tham chiếu của mô hình hiện tại, không phải kết quả chạy thí nghiệm mới khi soạn báo cáo. Các phép tính hai lớp ở chương 6 là ví dụ minh họa độc lập.')
table(['Chỉ số trên 1.490 mẫu kiểm tra','Giá trị'],[
['Accuracy','88,52%'],['Precision macro','88,41%'],['Recall macro','88,35%'],['F1 macro','88,33%'],['F1 weighted','88,49%']],[11.5,4.5])
p('Kết quả tương ứng 1.319 dự đoán đúng và 171 dự đoán sai. Ma trận nhầm lẫn có 25 mẫu sci.space được dự đoán thành rec.sport.baseball và 25 mẫu sci.space thành talk.politics.misc. Các con số xác định nơi cần khảo sát; muốn giải thích nguyên nhân phải đọc văn bản cụ thể, nhãn gốc và đặc trưng, không suy ra nguyên nhân chỉ từ số đếm.')
h('11.4 Điều kiện sử dụng và hướng phát triển')
p('Phương pháp phù hợp khi nhãn cố định, dữ liệu có nhãn đủ đại diện và nội dung có bằng chứng từ vựng liên quan chủ đề. Hạn chế gồm mất thứ tự từ, giả định đơn giản, lớp chồng lấn, từ mới và thay đổi miền. Không sử dụng kết quả như quyết định cuối cùng trong tác vụ cần hiểu ngữ cảnh sâu hoặc có hậu quả lớn.')
p('Nếu chuyển sang tiếng Việt, cần thu thập dữ liệu có nhãn tiếng Việt, thiết kế tách từ, học lại từ vựng và mô hình, rồi đánh giá trên tập kiểm tra độc lập. Nếu bổ sung lưu lịch sử lâu dài hoặc tải tệp, cần thiết kế lưu trữ và trích xuất nội dung riêng. Đây là định hướng nghiên cứu tiếp theo, chưa phải chức năng có trong giao diện hiện tại.')
boundary();old(14);boundary();old(15)

# Sửa phần dẫn nhập để phản ánh phạm vi đã bổ sung.
for par in doc.paragraphs:
    if par.text.startswith('Báo cáo gồm bảy chương'):
        par.text='Báo cáo gồm mười một chương. Chương 1 đến 7 giải thích cơ sở lý thuyết, ví dụ tính tay và giới hạn phương pháp. Chương 8 đến 11 trình bày ý tưởng hệ thống, chức năng, dữ liệu, logic xử lý, giao diện và thư viện theo hệ thống hiện có. Các số liệu hệ thống được ghi rõ là tham chiếu; báo cáo tập trung mô tả và nghiên cứu, không trình bày mã chương trình.'
    if par.text.startswith('Phạm vi chính là phân loại đơn nhãn'):
        par.text='Phạm vi chính là phân loại đơn nhãn theo chủ đề. Báo cáo dùng Bag of Words để diễn giải mô hình đa thức và TF-IDF để mô tả phương án hệ thống hiện có. Mô hình sử dụng bốn nhóm văn bản tiếng Anh của 20 Newsgroups. Phương pháp trình bày gồm tổng hợp tài liệu, diễn giải toán học, ví dụ tính tay và đối chiếu thiết kế với dữ liệu, dịch vụ phân loại và giao diện của dự án.'
    if par.text.startswith('Nhóm hệ thống hóa được quan hệ'):
        par.text='Nhóm hệ thống hóa quan hệ giữa tiền xử lý, biểu diễn véc-tơ và suy luận xác suất; xây dựng ví dụ tính tay và phân tích giới hạn phương pháp. Báo cáo còn mô tả ý tưởng hệ thống, chức năng, giao diện, dữ liệu bốn chủ đề tiếng Anh, logic học và dự đoán cùng vai trò các thư viện. Thiết kế TF-IDF kết hợp Multinomial Naive Bayes được đối chiếu với hệ thống hiện có, đáp ứng phạm vi nghiên cứu môn Trí tuệ nhân tạo.'

for ref in [
'[9] scikit-learn. The 20 newsgroups text dataset. https://scikit-learn.org/stable/datasets/real_world.html#the-20-newsgroups-text-dataset',
'[10] pandas. pandas.DataFrame. https://pandas.pydata.org/docs/reference/api/pandas.DataFrame.html',
'[11] Streamlit. Session State. https://docs.streamlit.io/develop/api-reference/caching-and-state/st.session_state',
'[12] joblib. joblib.load và joblib.dump. https://joblib.readthedocs.io/en/stable/generated/joblib.load.html',
'Tài liệu đối chiếu nội bộ: app.py; src/prepare_data.py; src/config.py; src/tfidf_pipeline.py; src/tune_alpha.py; src/classifier_service.py; requirements.txt và results/evaluation_summary.json. Các đường dẫn thuộc kho phan-loai-van-ban-naive-bayes của nhóm. Tài liệu trực tuyến được đối chiếu ngày 08/10/2026.'
]:p(ref)

# Tài liệu tham khảo được dàn gọn trong một trang, vẫn giữ toàn bộ nguồn.
in_refs=False
for par in doc.paragraphs:
    if par.text=='TÀI LIỆU THAM KHẢO':in_refs=True
    elif in_refs:
        par.alignment=WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.line_spacing=1.05
        par.paragraph_format.space_after=Pt(5)
        for run in par.runs:run.font.size=Pt(11.5)

assert len(doc.sections)==24, len(doc.sections)
for number,section in enumerate(doc.sections,1):
    section.different_first_page_header_footer=False
    section.footer.is_linked_to_previous=False
    for element in list(section.footer._element):section.footer._element.remove(element)
    par=section.footer.add_paragraph(str(number));par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_before=Pt(0);par.paragraph_format.space_after=Pt(0)
    for run in par.runs:run.font.name='Times New Roman';run.font.size=Pt(13);run.font.color.rgb=RGBColor(0,0,0)
    section.footer_distance=Cm(1)
    typ=section._sectPr.find(qn('w:pgNumType'))
    if typ is None:typ=OxmlElement('w:pgNumType');section._sectPr.append(typ)
    typ.set(qn('w:start'),str(number))
doc.core_properties.title='Nghiên cứu và trình bày một phương pháp phân loại văn bản bằng Multinomial Naive Bayes'
doc.core_properties.subject='Báo cáo môn Trí tuệ nhân tạo về lý thuyết và thiết kế hệ thống'
doc.core_properties.author='Nhóm 6'
doc.save(OUT)
print(OUT)

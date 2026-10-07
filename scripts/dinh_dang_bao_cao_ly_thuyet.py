"""Áp dụng bìa mẫu và mục lục, bảo toàn nội dung báo cáo lý thuyết AI."""
from copy import deepcopy
from pathlib import Path
import hashlib
import re
import shutil
from zipfile import ZipFile
from docx import Document
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).resolve().parents[1]
REF=Path('D:/Downloads/BCQLDAPM_Nhóm 11_06102026.docx')
OUT=ROOT/'docs/bao-cao-ly-thuyet-phan-loai-van-ban.docx'
ref_sha=hashlib.sha256(REF.read_bytes()).hexdigest()
ref=Document(REF)
doc=Document(OUT)
body=doc._element.body
intro=next(p for p in doc.paragraphs if p.text=='LỜI MỞ ĐẦU')
all_nodes=list(body)
intro_index=all_nodes.index(intro._p)
preserved_text=[''.join(e.itertext()) for e in all_nodes[intro_index:] if e.tag != qn('w:sectPr')]
for element in all_nodes[:intro_index]: body.remove(element)

logo=ROOT/'docs/assets/logo-ptit.png'
logo.parent.mkdir(parents=True,exist_ok=True)
logo.write_bytes(ZipFile(REF).read('word/media/image1.png'))
rid,_=doc.part.get_or_add_image(str(logo))
cover=[]
replace={9:'MÔN HỌC: TRÍ TUỆ NHÂN TẠO',
10:'ĐỀ TÀI: Nghiên cứu và trình bày một phương pháp\nphân loại văn bản bằng MULTINOMIAL NAIVE BAYES',
14:'Giảng viên hướng dẫn: ................................................',
15:'Thực hiện bởi nhóm sinh viên: Nhóm 6',
16:'1. Lại Huy Thịnh\tN23DVCN057',
17:'2. Nguyễn Hữu Đức\tN23DVCN012',
18:'3. Nguyễn Trần Mạnh Dũng\tN23DVCN015',
25:'TP.HCM, tháng 10/2026'}
for i, src in enumerate(ref.paragraphs[:26]):
    if i==2: continue
    node=deepcopy(src._p)
    par=Paragraph(node,doc._body)
    for blip in node.iter(qn('a:blip')): blip.set(qn('r:embed'),rid)
    if i in replace:
        if par.runs:
            par.runs[0].text=replace[i]
            for run in par.runs[1:]: run.text=''
        else: par.add_run(replace[i])
    if i in [16,17,18]:
        par.style=doc.styles['Normal']
        pr=node.get_or_add_pPr()
        for child in list(pr):
            if child.tag in [qn('w:numPr'),qn('w:ind'),qn('w:tabs')]: pr.remove(child)
        par.paragraph_format.left_indent=Cm(1.5)
        par.paragraph_format.tab_stops.add_tab_stop(Cm(10.8))
        par.paragraph_format.space_after=Pt(6)
        par.paragraph_format.line_spacing=1
    if i in [14,15]:
        par.alignment=WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.left_indent=Cm(0)
        par.paragraph_format.right_indent=Cm(0)
    if i==14:
        par.paragraph_format.space_before=Pt(38)
        par.paragraph_format.space_after=Pt(0)
    for run in par.runs:
        run.font.name='Times New Roman'
        if i==10:
            run.font.size=Pt(16)
            run.bold=True
        if i in [14,15,16,17,18,25]: run.font.size=Pt(13)
        run.font.color.rgb=RGBColor(0,0,0)
    if i==14:
        teacher_node=node
        continue
    cover.append(node)
    if i==15:
        x=deepcopy(node); xp=Paragraph(x,doc._body)
        xp.runs[0].text='Lớp: D23VHCNHT01-N'
        for run in xp.runs[1:]: run.text=''
        xp.paragraph_format.space_after=Pt(10)
        cover.append(x)
    if i==18:
        cover.append(teacher_node)

for index,node in enumerate(cover): body.insert(index,node)
anchor=intro
br=anchor.insert_paragraph_before();br.add_run().add_break(WD_BREAK.PAGE)
title=anchor.insert_paragraph_before('MỤC LỤC','Heading 1')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before=Pt(0)
title.paragraph_format.space_after=Pt(14)
title.runs[0].font.size=Pt(16)

current_page=2
toc=[]
for par in doc.paragraphs:
    if par.text=='LỜI MỞ ĐẦU': current_page=3
    elif current_page>=3 and par.style.name=='Heading 1': current_page+=1
    if current_page<3: continue
    if par.style.name=='Heading 1' or (par.style.name=='Heading 2' and re.match(r'^\d+\.\d+\s',par.text)):
        text=par.text
        is_chapter=not re.match(r'^\d+\.',text)
        if text=='MỤC LỤC': continue
        # Các tiêu đề trang 2.2, 3.2, 5.2, 7.2 cũng là mục con của chương.
        if par.style.name=='Heading 2' and toc and toc[-1][0].casefold()==text.casefold(): continue
        toc.append((text,current_page,is_chapter))
pdfmetrics.registerFont(TTFont('TocRegular','C:/Windows/Fonts/times.ttf'))
pdfmetrics.registerFont(TTFont('TocBold','C:/Windows/Fonts/timesbd.ttf'))
# Dấu chấm là ký tự thật; cột số trang cố định, không phụ thuộc hỗ trợ tab leader.
toc_table=doc.add_table(rows=0,cols=2)
toc_table.autofit=False
toc_table.columns[0].width=Cm(15.3)
toc_table.columns[1].width=Cm(0.7)
tbl_pr=toc_table._tbl.tblPr
borders=OxmlElement('w:tblBorders')
for side in ['top','left','bottom','right','insideH','insideV']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'nil');borders.append(e)
tbl_pr.append(borders)
for text,num,major in toc:
    row=toc_table.add_row()
    no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
    for i,cell in enumerate(row.cells):
        cell.width=Cm(15.3 if i==0 else 0.7)
        margins=OxmlElement('w:tcMar')
        for side in ['top','left','bottom','right']:
            e=OxmlElement('w:'+side);e.set(qn('w:w'),'0');e.set(qn('w:type'),'dxa');margins.append(e)
        cell._tc.get_or_add_tcPr().append(margins)
        x=cell.paragraphs[0]
        x.paragraph_format.line_spacing=Pt(17)
        x.paragraph_format.space_before=Pt(0)
        x.paragraph_format.space_after=Pt(7)
        x.paragraph_format.keep_with_next=False
        x.paragraph_format.left_indent=Cm(0.45 if i==0 and not major else 0)
        x.paragraph_format.right_indent=Cm(0)
        x.paragraph_format.first_line_indent=Cm(0)
        x.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.RIGHT
        if i==0:
            font='TocBold' if major else 'TocRegular'
            available=Cm(15.3).pt-(0 if major else Cm(0.45).pt)-4
            dots=max(3,int((available-pdfmetrics.stringWidth(text+' ',font,13))/pdfmetrics.stringWidth('.',font,13)))
            content=text+' '+'.'*dots
        else: content=str(num)
        r=x.add_run(content)
        r.font.name='Times New Roman';r.font.size=Pt(13)
        r.font.color.rgb=RGBColor(0,0,0);r.bold=major
intro._p.addprevious(toc_table._tbl)
br=anchor.insert_paragraph_before();br.add_run().add_break(WD_BREAK.PAGE)

for par in doc.paragraphs:
    for run in par.runs:
        run.font.color.rgb=RGBColor(0,0,0)
for style in doc.styles:
    if style.type in [1,2]: style.font.color.rgb=RGBColor(0,0,0)
assert preserved_text==[''.join(e.itertext()) for e in list(body)[list(body).index(intro._p):] if e.tag!=qn('w:sectPr')]
doc.core_properties.title='Nghiên cứu và trình bày một phương pháp phân loại văn bản bằng MULTINOMIAL NAIVE BAYES'
doc.save(OUT)
shutil.copy2(OUT,ROOT/'docs/bao-cao-tri-tue-nhan-tao-nhom-6.docx')
assert hashlib.sha256(REF.read_bytes()).hexdigest()==ref_sha
print('Reference SHA256',ref_sha)
print('TOC entries',len(toc))
print('Preserved all theory paragraphs and tables')
